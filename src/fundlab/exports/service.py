"""中文 XLSX/CSV 与来源元数据；先生成完整文件，再发布并在失败时回滚。"""
import csv
from io import BytesIO, StringIO
from decimal import Decimal, localcontext
import json
import os
from pathlib import Path
import re
import shutil
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from fundlab.config import AppConfig
from fundlab.errors import FundlabError
from fundlab.models import Fund, utc_now
from fundlab.parsers.history import FIELD_DEFINITIONS, RATIO_FIELDS, canonical_decimal
from fundlab.validation import validate_code


def safe_text(value: str) -> str:
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def fund_folder(fund: Fund) -> str:
    validate_code(fund.code)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", fund.name).strip(" .")[:60].rstrip(" .") or "未命名基金"
    return f"{fund.code}_{name}"


def publish_files(staged: Path, target: Path, names: tuple[str, ...]):
    """单个文件用原子替换；组发布失败则恢复已经替换的文件。"""
    backups = staged / "backups"
    backups.mkdir()
    existed = {}
    for name in names:
        destination = target / name
        existed[name] = destination.exists()
        if existed[name]:
            shutil.copy2(destination, backups / name)
    try:
        for name in names:
            os.replace(staged / name, target / name)
    except (OSError, KeyboardInterrupt):
        try:
            # os.replace 可能已完成，但 Ctrl+C 在下一条 Python 指令前到达。
            # 源临时文件是否仍存在能识别实际替换结果，不依赖另行 append 的时机。
            for name in reversed(names):
                if (staged / name).exists():
                    continue
                if existed[name]:
                    os.replace(backups / name, target / name)
                else:
                    (target / name).unlink(missing_ok=True)
        except (OSError, KeyboardInterrupt):
            raise FundlabError("EXPORT_RECOVERY_REQUIRED", f"导出发布和回滚未全部完成，备份保留在 {backups}；数据库未丢失。") from None
        raise


def _write_csv(path: Path, rows: list[dict]):
    path.write_bytes(render_csv(rows))


def render_csv(rows: list[dict]) -> bytes:
    """共享既有十列和精度规则；供文件导出及只读网页下载使用。"""
    with StringIO(newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([label for _, _, label, _ in FIELD_DEFINITIONS])
        for row in rows:
            values = []
            for key, _, _, _ in FIELD_DEFINITIONS:
                value = row[key]
                if value is None:
                    values.append("")
                elif key in RATIO_FIELDS:
                    with localcontext() as ctx:
                        ctx.prec = 60
                        values.append(canonical_decimal(Decimal(value) * 100) + "%")
                else:
                    # 非比例列只有标准日期或已经严格解析的十进制数。
                    values.append(value)
            writer.writerow(values)
        return handle.getvalue().encode("utf-8-sig")


def _write_xlsx(path: Path | BytesIO, fund: Fund, rows: list[dict], metadata: dict):
    book = Workbook()
    sheet = book.active
    sheet.title = "历史数据"
    sheet.append([label for _, _, label, _ in FIELD_DEFINITIONS])
    for row in rows:
        values = []
        for key, _, _, _ in FIELD_DEFINITIONS:
            value = row[key]
            if value is None or key in ("price_date", "nav_date"):
                values.append(value)
            else:
                number = Decimal(value)
                # Excel 超过 15 位有效数字会丢精度，此时保留为文本。
                values.append(number if len(number.as_tuple().digits) <= 15 else value)
        sheet.append(values)
    formats = {"close": "0.000", "nav": "0.0000", "source_premium": "0.00%",
               "turnover": "#,##0.00", "shares": "#,##0.####", "shares_added": "#,##0.####",
               "shares_change": "0.000%", "index_change": "0.00%"}
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="2463A6")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for column, (key, _, _, _) in enumerate(FIELD_DEFINITIONS, start=1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 19
        for cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            cell = cells[0]
            if cell.data_type == "n":
                cell.number_format = formats.get(key, "General")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="EEF4FB")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    note = book.create_sheet("说明")
    info = [("基金代码", fund.code), ("基金名称", fund.name), ("类别", fund.category),
            ("来源", fund.source_url), ("导出时间(UTC)", metadata["exported_at"]),
            ("导出行数" if metadata.get("selection") else "本地累计行数", str(len(rows))),
            ("导出起始日期" if metadata.get("selection") else "本地起始日期", rows[-1]["price_date"]),
            ("导出截至日期" if metadata.get("selection") else "本地截至日期", rows[0]["price_date"]),
            ("当前源窗口", json.dumps(metadata["coverage"], ensure_ascii=False)),
            ("单位", "成交额：万元；份额及新增：万份；比例列为百分比显示"),
            ("日期口径", "价格日期与净值日期分开保存；不承诺两者是同步报价"),
            ("完整性", "当前可访问窗口加本地累积，不代表成立以来全部历史"),
            ("精度", "数据库保留十进制字符串及原值；Excel超过15位有效数字的数值转文本保存")]
    for label, value in info:
        note.append([label, safe_text(value)])
    note.column_dimensions["A"].width = 24
    note.column_dimensions["B"].width = 90
    for cells in note:
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    try:
        book.save(path)
    finally:
        book.close()


def render_history_file(fund: Fund, rows: list[dict], coverage: dict, format: str) -> bytes:
    """从只读快照生成下载，不发布归档、不创建运行记录、不写数据库。"""
    if not rows:
        raise FundlabError("NO_DATA", "当前日期区间没有记录，请调整区间后下载。")
    if format == "csv":
        return render_csv(rows)
    if format != "xlsx":
        raise FundlabError("VALIDATION_ERROR", "仅支持 Excel 或 CSV 下载。")
    with BytesIO() as stream:
        _write_xlsx(stream, fund, rows, {"exported_at": utc_now(), "coverage": coverage, "selection": True})
        return stream.getvalue()


def export_history(config: AppConfig, fund: Fund, rows: list[dict], coverage: dict, run_id: str) -> list[dict]:
    if not rows:
        raise FundlabError("NO_DATA", "本地没有该基金的历史记录，未导出空文件。")
    config.funds_dir.mkdir(parents=True, exist_ok=True)
    target = config.funds_dir / fund_folder(fund)
    if not target.resolve().is_relative_to(config.funds_dir.resolve()):
        raise FundlabError("VALIDATION_ERROR", "基金导出目录超出配置范围。")
    target.mkdir(exist_ok=True)
    names = ("history.csv", "history.xlsx", "metadata.json")
    metadata = {"schema_version": 1, "fund": {"code": fund.code, "name": fund.name, "category": fund.category,
                "source": fund.source, "source_url": fund.source_url}, "run_id": run_id, "rows": len(rows),
                "local_start": rows[-1]["price_date"], "local_end": rows[0]["price_date"],
                "exported_at": utc_now(), "coverage": coverage,
                "field_definitions": [{"key": key, "source_key": source, "label": label, "unit": unit}
                                      for key, source, label, unit in FIELD_DEFINITIONS]}
    staged = Path(tempfile.mkdtemp(prefix=".export-", dir=target))
    retain_backups = False
    try:
        _write_csv(staged / names[0], rows)
        _write_xlsx(staged / names[1], fund, rows, metadata)
        (staged / names[2]).write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
        publish_files(staged, target, names)
    except FundlabError as error:
        retain_backups = error.code == "EXPORT_RECOVERY_REQUIRED"
        raise
    except OSError:
        raise FundlabError("EXPORT_ERROR", "数据库数据仍保留；导出未完成，请关闭占用的 Excel/CSV 文件并检查磁盘后重试。") from None
    finally:
        if not retain_backups and staged.resolve().is_relative_to(target.resolve()) and staged.name.startswith(".export-"):
            shutil.rmtree(staged, ignore_errors=True)
    return [{"format": name.rsplit(".", 1)[-1], "path": str(target / name), "rows": len(rows), "status": "succeeded"}
            for name in names]
