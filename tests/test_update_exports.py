"""手动更新与导出的关键保证；仅临时目录、明确标注的人工输入。"""
import csv
import ctypes
from ctypes import wintypes
import json
import os
import subprocess
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from fundlab.config import AppConfig
from fundlab.errors import FundlabError
from fundlab.exports.service import export_history, fund_folder
from fundlab.models import Fund, SourceWindow
from fundlab.services.update import retry_export, update_one
from fundlab.storage.repository import Repository
from test_data_storage import fixture_row, FUND, STAMP


class UpdateExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.config = AppConfig(Path(self.temp.name))
        silent_progress = patch("fundlab.services.update.print")
        silent_progress.start()
        self.addCleanup(silent_progress.stop)

    def window(self, rows=None):
        return SourceWindow(FUND, rows if rows is not None else [fixture_row()], STAMP,
                            {"kind": "source_window", "full_history_verified": False})

    def update(self, window=None):
        with patch("fundlab.services.update.fetch_window", return_value=window or self.window()):
            return update_one(self.config, FUND.code)

    def test_first_and_repeat_update_and_export_values(self):
        first = self.update()
        second = self.update()
        self.assertEqual((first.status, first.inserted, second.status, second.unchanged), ("succeeded", 1, "succeeded", 1))
        folder = self.config.funds_dir / fund_folder(FUND)
        self.assertTrue((folder / "history.csv").read_bytes().startswith(b"\xef\xbb\xbf"))
        with (folder / "history.csv").open(encoding="utf-8-sig", newline="") as handle:
            data = list(csv.reader(handle))
        self.assertEqual((len(data), len(data[0]), data[1][4], data[1][9]), (2, 10, "19.96%", "-2.5%"))
        book = load_workbook(folder / "history.xlsx", data_only=False)
        try:
            self.assertEqual(book["历史数据"].max_row, 2)
            self.assertAlmostEqual(book["历史数据"]["E2"].value, 0.1996)
            self.assertEqual(book["历史数据"]["E2"].number_format, "0.00%")
            self.assertEqual(book["说明"]["B1"].value, "000001")
        finally:
            book.close()

    def test_failed_export_preserves_database_and_retries_without_network(self):
        with patch("fundlab.services.update.fetch_window", return_value=self.window()), patch(
                "fundlab.services.update.export_history", side_effect=FundlabError("EXPORT_ERROR", "测试导出失败")):
            failed = update_one(self.config, FUND.code)
        self.assertEqual(failed.status, "export_failed")
        self.assertEqual(len(Repository(self.config.database_path).query_history(FUND.code)), 1)
        with patch("fundlab.services.update.fetch_window", side_effect=AssertionError("不可联网")):
            result = retry_export(self.config, FUND.code)
        self.assertEqual(result.status, "succeeded")

    def test_parse_failure_leaves_old_history_and_records_failure(self):
        self.update()
        broken = fixture_row(price="NaN")
        result = self.update(self.window([broken]))
        repo = Repository(self.config.database_path)
        self.assertEqual(result.status, "failed")
        self.assertEqual(repo.query_history(FUND.code)[0]["close"], "2.345")
        self.assertEqual(repo.list_runs(1)[0]["status"], "failed")

    def test_publish_failure_rolls_back_entire_export_set(self):
        self.update()
        folder = self.config.funds_dir / fund_folder(FUND)
        before = {name: (folder / name).read_bytes() for name in ("history.csv", "history.xlsx", "metadata.json")}
        real_replace = os.replace
        def fail_xlsx(source, destination):
            if Path(source).name == "history.xlsx" and Path(source).parent.name != "backups":
                raise PermissionError("测试模拟 Excel 占用")
            return real_replace(source, destination)
        with patch("fundlab.exports.service.os.replace", side_effect=fail_xlsx):
            result = self.update(self.window([fixture_row(price="2.5")]))
        self.assertEqual(result.status, "export_failed")
        self.assertEqual(Repository(self.config.database_path).query_history(FUND.code)[0]["close"], "2.5")
        for name, content in before.items():
            self.assertEqual((folder / name).read_bytes(), content)

    def test_interrupt_during_publish_restores_previous_export_set(self):
        self.update()
        folder = self.config.funds_dir / fund_folder(FUND)
        before = {name: (folder / name).read_bytes() for name in ("history.csv", "history.xlsx", "metadata.json")}
        real_replace = os.replace
        def interrupt_xlsx(source, destination):
            if Path(source).name == "history.xlsx" and Path(source).parent.name != "backups":
                raise KeyboardInterrupt()
            return real_replace(source, destination)
        with patch("fundlab.exports.service.os.replace", side_effect=interrupt_xlsx), self.assertRaises(KeyboardInterrupt):
            self.update(self.window([fixture_row(price="2.5")]))
        for name, content in before.items():
            self.assertEqual((folder / name).read_bytes(), content)
        repo = Repository(self.config.database_path)
        self.assertEqual(repo.query_history(FUND.code)[0]["close"], "2.5")
        self.assertEqual(repo.list_runs(1)[0]["status"], "interrupted")

    def test_failed_rollback_retains_backups_for_recovery(self):
        self.update()
        folder = self.config.funds_dir / fund_folder(FUND)
        old_csv = (folder / "history.csv").read_bytes()
        real_replace = os.replace
        def fail_publish_and_rollback(source, destination):
            source = Path(source)
            if source.name == "history.xlsx" or source.parent.name == "backups":
                raise PermissionError("测试发布及回滚失败")
            return real_replace(source, destination)
        with patch("fundlab.exports.service.os.replace", side_effect=fail_publish_and_rollback):
            result = self.update(self.window([fixture_row(price="2.5")]))
        self.assertEqual((result.status, result.error_code), ("export_failed", "EXPORT_RECOVERY_REQUIRED"))
        backups = list(folder.glob(".export-*/backups/history.csv"))
        self.assertEqual(len(backups), 1)
        self.assertEqual(backups[0].read_bytes(), old_csv)

    def test_interrupt_immediately_after_replace_still_restores_old_file(self):
        self.update()
        folder = self.config.funds_dir / fund_folder(FUND)
        before = {name: (folder / name).read_bytes() for name in ("history.csv", "history.xlsx", "metadata.json")}
        real_replace = os.replace
        def replace_then_interrupt(source, destination):
            result = real_replace(source, destination)
            if Path(source).name == "history.csv" and Path(source).parent.name != "backups":
                raise KeyboardInterrupt()
            return result
        with patch("fundlab.exports.service.os.replace", side_effect=replace_then_interrupt), self.assertRaises(KeyboardInterrupt):
            self.update(self.window([fixture_row(price="2.5")]))
        for name, content in before.items():
            self.assertEqual((folder / name).read_bytes(), content)

    def test_formula_text_and_windows_path_are_safe(self):
        special = Fund("000001", "=1+1/中文:*?", "QDII", FUND.source_url)
        source = self.window()
        source = SourceWindow(special, source.rows, source.fetched_at, source.coverage)
        result = self.update(source)
        self.assertEqual(result.status, "succeeded")
        folder = self.config.funds_dir / fund_folder(special)
        self.assertTrue(folder.resolve().is_relative_to(self.config.funds_dir.resolve()))
        book = load_workbook(folder / "history.xlsx", data_only=False)
        try:
            cell = book["说明"]["B2"]
            self.assertEqual(cell.data_type, "s")
            self.assertTrue(cell.value.startswith("'="))
        finally:
            book.close()

    def test_csv_preserves_high_precision_ratio(self):
        source_percent = "19.961234567890123456789012345678901234"
        result = self.update(self.window([fixture_row(discount_rt=source_percent)]))
        self.assertEqual(result.status, "succeeded")
        path = self.config.funds_dir / fund_folder(FUND) / "history.csv"
        with path.open(encoding="utf-8-sig", newline="") as handle:
            data = list(csv.reader(handle))
        self.assertEqual(data[1][4], source_percent + "%")

    def test_unknown_formula_fund_code_cannot_escape_folder(self):
        with self.assertRaises(FundlabError):
            fund_folder(Fund("../bad", "名字", "QDII", FUND.source_url))

    @unittest.skipUnless(os.name == "nt", "Windows 文件共享锁行为")
    def test_real_windows_file_lock_preserves_old_export(self):
        self.update()
        folder = self.config.funds_dir / fund_folder(FUND)
        filename = folder / "history.xlsx"
        content = filename.read_bytes()
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                      wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        kernel.CreateFileW.restype = wintypes.HANDLE
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        handle = kernel.CreateFileW(str(filename), 0x80000000, 0, None, 3, 0x80, None)
        self.assertNotEqual(handle, wintypes.HANDLE(-1).value)
        try:
            result = self.update()
            self.assertEqual(result.status, "export_failed")
        finally:
            kernel.CloseHandle(handle)
        self.assertEqual(filename.read_bytes(), content)

    @unittest.skipUnless(os.name == "nt", "Windows CMD 启动入口")
    def test_windows_launcher_uses_crlf_and_rejects_blank_without_collection(self):
        root = AppConfig.default().root
        for name in ("login.cmd", "update.cmd"):
            content = (root / "scripts" / name).read_bytes()
            self.assertEqual(content.count(b"\n"), content.count(b"\r\n"))
            self.assertFalse(content.startswith(b"\xef\xbb\xbf"))
        database = root / "data" / "fundlab.sqlite3"
        before = database.stat().st_mtime_ns if database.exists() else None
        process = subprocess.run([str(Path(os.environ.get("SystemRoot", r"C:\Windows")) / "System32/cmd.exe"),
                                  "/d", "/c", "scripts\\update.cmd"], cwd=root, input="\n\n",
                                 capture_output=True, encoding="utf-8", timeout=20,
                                 creationflags=subprocess.CREATE_NO_WINDOW)
        self.assertEqual(process.returncode, 1)
        output = process.stdout + process.stderr
        self.assertIn("VALIDATION_ERROR", output)
        self.assertNotIn("not recognized", output)
        self.assertEqual(before, database.stat().st_mtime_ns if database.exists() else None)


if __name__ == "__main__":
    unittest.main()
