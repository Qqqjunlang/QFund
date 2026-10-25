"""只读网页查询、下载内容与HTTP边界；fixture仅供测试。"""
import csv
from datetime import date, timedelta
from io import BytesIO, StringIO
import json
from pathlib import Path
import tempfile
from threading import Thread
import unittest
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from fundlab.config import AppConfig
from fundlab.errors import FundlabError
from fundlab.services.data_query import download_history, history_page, overview
from fundlab.parsers.history import parse_history
from fundlab.storage.repository import Repository
from fundlab.web.server import make_server
from test_data_storage import FUND, STAMP, fixture_row


class WebQueryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.config = AppConfig(Path(self.temp.name))
        self.repo = Repository(self.config.database_path)
        self.repo.start_run("test-query", FUND.code, "initial")
        rows = [fixture_row((date(2020, 1, 2) + timedelta(days=i)).isoformat()) for i in range(51)]
        rows[-1] = fixture_row(rows[-1]["id"], discount_rt="-", amount_incr="0")
        self.repo.commit_window(FUND, parse_history(rows, FUND.code, STAMP), "test-query", {"kind": "source_window"})
        self.repo.finish_run({"run_id": "test-query", "code": FUND.code, "status": "succeeded", "inserted": 51})

    def test_overview_ignores_export_runs_and_keeps_leading_zero(self):
        self.repo.start_run("test-export", FUND.code, "export")
        result = overview(self.config)
        self.assertEqual(result["funds"][0]["code"], "000001")
        self.assertEqual(result["funds"][0]["records"], 51)
        self.assertEqual(result["funds"][0]["fetched_at"], STAMP)
        self.assertEqual(result["latest"]["code"], FUND.code)
        self.assertNotIn("result_json", result["latest"])

    def test_missing_database_empty_state_does_not_create_files(self):
        config = AppConfig(self.config.root / "empty")
        self.assertEqual(overview(config)["funds"], [])
        with self.assertRaises(FundlabError):
            history_page(config, FUND.code)
        self.assertFalse(config.root.exists())

    def test_page_sort_missing_zero_and_percentage_units(self):
        result = history_page(self.config, FUND.code)
        self.assertEqual((result["total"], len(result["rows"]), result["pages"]), (51, 50, 2))
        self.assertIsNone(result["rows"][0]["source_premium"])
        self.assertEqual(result["rows"][0]["shares_added"], "0")
        self.assertEqual(result["rows"][1]["source_premium"], "19.96%")
        self.assertEqual(result["rows"][1]["index_change"], "-2.5%")
        self.assertEqual(history_page(self.config, FUND.code, page=2)["rows"][0]["price_date"], "2020-01-02")

    def test_inclusive_dates_empty_and_invalid_range(self):
        result = history_page(self.config, FUND.code, "2020-01-02", "2020-01-03")
        self.assertEqual(result["matched"], 2)
        self.assertEqual(history_page(self.config, FUND.code, "2025-01-01")["matched"], 0)
        for start, end in (("2020-02-30", ""), ("20200101", ""), ("2020-02-02", "2020-01-01")):
            with self.subTest(start=start), self.assertRaises(FundlabError):
                history_page(self.config, FUND.code, start, end)
        with self.assertRaises(FundlabError):
            history_page(self.config, FUND.code, page=0)

    def test_download_includes_all_filtered_rows_not_just_page(self):
        _, content = download_history(self.config, FUND.code, "csv")
        rows = list(csv.reader(StringIO(content.decode("utf-8-sig"))))
        self.assertEqual(len(rows), 52)
        self.assertEqual((rows[1][4], rows[2][4], rows[1][7]), ("", "19.96%", "0"))
        name, content = download_history(self.config, FUND.code, "csv", "2020-01-02", "2020-01-03")
        self.assertTrue(name.startswith("000001_history_2020-01-02_2020-01-03"))
        self.assertEqual(len(list(csv.reader(StringIO(content.decode("utf-8-sig"))))), 3)

    def test_xlsx_numeric_units_code_and_scope(self):
        _, content = download_history(self.config, FUND.code, "xlsx", "2020-01-02", "2020-01-03")
        book = load_workbook(BytesIO(content))
        try:
            self.assertEqual(book["历史数据"].max_row, 3)
            self.assertAlmostEqual(book["历史数据"]["E2"].value, .1996)
            self.assertEqual(book["历史数据"]["E2"].number_format, "0.00%")
            self.assertEqual(book["说明"]["B1"].value, "000001")
            self.assertEqual(book["说明"]["A6"].value, "导出行数")
        finally:
            book.close()

    def test_download_empty_or_invalid_format_fails(self):
        for format, start in (("csv", "2025-01-01"), ("html", "")):
            with self.subTest(format=format), self.assertRaises(FundlabError):
                download_history(self.config, FUND.code, format, start)

    def test_queries_and_downloads_leave_database_and_archive_untouched(self):
        before = self.config.database_path.read_bytes()
        runs = self.repo.list_runs()
        overview(self.config)
        history_page(self.config, FUND.code)
        for format in ("xlsx", "csv"):
            download_history(self.config, FUND.code, format)
        self.assertEqual(self.config.database_path.read_bytes(), before)
        self.assertEqual(self.repo.list_runs(), runs)
        self.assertFalse(self.config.funds_dir.exists())

    def start_server(self):
        server = make_server(self.config, 0)
        thread = Thread(target=server.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        return f"http://127.0.0.1:{server.server_port}"

    def test_http_assets_query_and_attachment(self):
        base = self.start_server()
        with urlopen(base + "/") as response:
            self.assertIn("数据查询".encode(), response.read())
        with urlopen(base + "/api/history?code=000001&start=2020-01-02&end=2020-01-03") as response:
            self.assertEqual(json.load(response)["matched"], 2)
            self.assertEqual(response.headers["Cache-Control"], "no-store")
        with urlopen(base + "/api/download?code=000001&format=csv") as response:
            self.assertIn("attachment", response.headers["Content-Disposition"])
            self.assertTrue(response.read().startswith(b"\xef\xbb\xbf"))

    def test_http_rejects_file_paths_host_and_invalid_queries(self):
        base = self.start_server()
        for path, status in (("/.local/edge-profile", 404), ("/../data/fundlab.sqlite3", 404),
                             ("/api/history?code=000001&page=NaN", 400), ("/api/history?code=000001&code=513100", 400),
                             ("/api/history?code=999999", 404), ("/api/download?code=000001&format=html", 400)):
            with self.subTest(path=path), self.assertRaises(HTTPError) as caught:
                urlopen(base + path)
            self.assertEqual(caught.exception.code, status)
        with self.assertRaises(HTTPError) as caught:
            urlopen(Request(base + "/api/overview", headers={"Host": "example.invalid"}))
        self.assertEqual(caught.exception.code, 403)


if __name__ == "__main__":
    unittest.main()
